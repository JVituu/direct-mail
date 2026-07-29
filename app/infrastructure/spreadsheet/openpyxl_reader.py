from collections.abc import Iterable
from datetime import date, datetime

from openpyxl import load_workbook

from app.domain.value_objects.spreadsheet_row import SpreadsheetRow


class OpenpyxlSpreadsheetReader:
    def list_sheets(self, file_path: str) -> list[str]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def read_headers(self, file_path: str, sheet_name: str) -> list[str]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration:
                return []
            return self._normalize_headers(raw_headers)
        finally:
            workbook.close()

    def iter_rows(self, file_path: str, sheet_name: str) -> Iterable[SpreadsheetRow]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration:
                return

            columns = self._normalize_headers(raw_headers)
            for row_number, raw_values in enumerate(rows, start=2):
                values = self._row_to_dict(columns, raw_values)
                if any(value for value in values.values()):
                    yield SpreadsheetRow(row_number=row_number, values=values)
        finally:
            workbook.close()

    def _row_to_dict(
        self,
        columns: list[str],
        raw_values: tuple[object, ...],
    ) -> dict[str, str]:
        values: dict[str, str] = {}
        for index, column in enumerate(columns):
            raw_value = raw_values[index] if index < len(raw_values) else None
            values[column] = self._to_text(raw_value)
        return values

    def _normalize_headers(self, raw_headers: tuple[object, ...]) -> list[str]:
        headers: list[str] = []
        seen: dict[str, int] = {}

        for index, raw_header in enumerate(raw_headers, start=1):
            header = self._to_text(raw_header) or f"Coluna {index}"
            count = seen.get(header, 0) + 1
            seen[header] = count
            if count > 1:
                header = f"{header} {count}"
            headers.append(header)

        return headers

    def _to_text(self, value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(value, date):
            return value.isoformat()
        return str(value).strip()
