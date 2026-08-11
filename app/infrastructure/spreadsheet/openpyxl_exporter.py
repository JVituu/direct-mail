from collections.abc import Iterable, Sequence
import hashlib
from pathlib import Path
import re
from unicodedata import combining, normalize

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font as ExcelFont, PatternFill, Side
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QMarginsF, QRectF, Qt
from PySide6.QtGui import (
    QColor,
    QFont,
    QFontMetrics,
    QGuiApplication,
    QPageLayout,
    QPageSize,
    QPainter,
    QPdfWriter,
    QPen,
)

from app.domain.entities.guest_record import GuestRecord


class OpenpyxlSelectedGuestsExporter:
    _qt_app: QGuiApplication | None = None
    _header_fill = PatternFill("solid", fgColor="245783")
    _header_font = ExcelFont(color="FFFFFF", bold=True)
    _header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _body_alignment = Alignment(vertical="center", wrap_text=True)
    _thin_side = Side(style="thin", color="DDE4EF")
    _cell_border = Border(left=_thin_side, right=_thin_side, top=_thin_side, bottom=_thin_side)
    _alternate_fill = PatternFill("solid", fgColor="F1F6FD")
    _white_fill = PatternFill("solid", fgColor="FFFFFF")
    _category_fills = (
        PatternFill("solid", fgColor="DBEAFE"),
        PatternFill("solid", fgColor="DCFCE7"),
        PatternFill("solid", fgColor="FEF3C7"),
        PatternFill("solid", fgColor="EDE9FE"),
        PatternFill("solid", fgColor="FAE8FF"),
        PatternFill("solid", fgColor="CFFAFE"),
        PatternFill("solid", fgColor="FFE4E6"),
        PatternFill("solid", fgColor="E0F2FE"),
        PatternFill("solid", fgColor="F3E8FF"),
        PatternFill("solid", fgColor="CCFBF1"),
        PatternFill("solid", fgColor="FFEDD5"),
        PatternFill("solid", fgColor="E2E8F0"),
    )
    _email_column_words = ("email", "e-mail", "mail")
    _category_column_words = ("categoria", "category")

    def export(
        self,
        output_path: str,
        columns: Sequence[str],
        guests: Iterable[GuestRecord],
        sheet_name: str = "Selecionados",
    ) -> int:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        worksheet = self._replace_default_sheet(
            workbook,
            self._safe_sheet_name(sheet_name, set()),
        )
        count = self._write_table(worksheet, columns, guests)

        workbook.save(path)
        workbook.close()
        return count

    def export_by_category(
        self,
        output_path: str,
        columns: Sequence[str],
        guests: Iterable[GuestRecord],
        category_column: str,
    ) -> int:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        grouped_guests: dict[str, list[GuestRecord]] = {}
        for guest in guests:
            category = self._category_value(guest.data.get(category_column, ""))
            grouped_guests.setdefault(category, []).append(guest)

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        used_sheet_names = {"Resumo"}
        summary = workbook.create_sheet("Resumo")

        total = 0
        summary_rows: list[dict[str, str]] = []
        for category in sorted(grouped_guests, key=str.casefold):
            category_guests = grouped_guests[category]
            summary_rows.append(
                {
                    "Categoria": category,
                    "Quantidade": len(category_guests),
                }
            )
            total += len(category_guests)
        self._write_table(summary, ("Categoria", "Quantidade"), self._guest_dicts(summary_rows))

        for category in sorted(grouped_guests, key=str.casefold):
            worksheet = workbook.create_sheet(self._safe_sheet_name(category, used_sheet_names))
            self._write_table(worksheet, columns, grouped_guests[category])

        workbook.save(path)
        workbook.close()
        return total

    def _replace_default_sheet(self, workbook: Workbook, title: str):
        worksheet = workbook.active
        worksheet.title = title
        return worksheet

    def _write_table(
        self,
        worksheet: object,
        columns: Sequence[str],
        guests: Iterable[GuestRecord],
    ) -> int:
        clean_columns = tuple(str(column) for column in columns)
        display_columns = [self._format_header_text(column) for column in clean_columns]
        worksheet.append(display_columns)
        self._style_header_row(worksheet, len(clean_columns))

        widths = [len(column) for column in display_columns]
        category_indexes = self._category_column_indexes(clean_columns)
        count = 0

        for guest in guests:
            row_values = [
                self._format_xlsx_value(column, guest.data.get(column, ""))
                for column in clean_columns
            ]
            worksheet.append(row_values)
            count += 1
            row_number = count + 1
            self._style_body_row(
                worksheet,
                row_number,
                row_values,
                category_indexes,
                widths,
            )

        self._finish_worksheet_layout(worksheet, len(clean_columns), count, widths)
        return count

    def _style_header_row(self, worksheet: object, column_count: int) -> None:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.row_dimensions[1].height = 24
        for column_index in range(1, column_count + 1):
            cell = worksheet.cell(row=1, column=column_index)
            cell.fill = self._header_fill
            cell.font = self._header_font
            cell.alignment = self._header_alignment
            cell.border = self._cell_border

    def _style_body_row(
        self,
        worksheet: object,
        row_number: int,
        values: Sequence[object],
        category_indexes: set[int],
        widths: list[int],
    ) -> None:
        row_fill = self._alternate_fill if row_number % 2 == 1 else self._white_fill
        max_lines = 1
        for column_index, value in enumerate(values, start=1):
            text = "" if value is None else str(value)
            max_lines = max(max_lines, text.count("\n") + 1)
            widths[column_index - 1] = max(
                widths[column_index - 1],
                self._display_width(text),
            )

            cell = worksheet.cell(row=row_number, column=column_index)
            cell.fill = self._category_fill(text) if column_index in category_indexes and text.strip() else row_fill
            cell.alignment = self._body_alignment
            cell.border = self._cell_border

        worksheet.row_dimensions[row_number].height = max(21, max_lines * 18)

    def _finish_worksheet_layout(
        self,
        worksheet: object,
        column_count: int,
        row_count: int,
        widths: Sequence[int],
    ) -> None:
        if column_count <= 0:
            return

        last_column_letter = get_column_letter(column_count)
        worksheet.auto_filter.ref = f"A1:{last_column_letter}{max(row_count + 1, 1)}"
        for column_index, width in enumerate(widths, start=1):
            column_letter = get_column_letter(column_index)
            worksheet.column_dimensions[column_letter].width = min(max(width + 3, 12), 48)

    def _guest_dicts(self, rows: Iterable[dict[str, object]]) -> Iterable[GuestRecord]:
        for index, row in enumerate(rows, start=1):
            yield GuestRecord(
                id=index,
                import_id=0,
                sheet_name="Resumo",
                row_number=index,
                verification_code="",
                data=row,
                selected=False,
            )

    def _format_header_text(self, value: object) -> str:
        text = str(value or "").strip().casefold()
        if not text:
            return ""
        return re.sub(
            r"(^|[\s/])([^\W\d_])",
            lambda match: f"{match.group(1)}{match.group(2).upper()}",
            text,
        )

    def _format_xlsx_value(self, column_name: str, value: object) -> object:
        if isinstance(value, (int, float)):
            return value
        text = "" if value is None else str(value).strip()
        if self._is_email_column(column_name):
            return text.casefold()
        return text.upper()

    def _category_column_indexes(self, columns: Sequence[str]) -> set[int]:
        return {
            index
            for index, column in enumerate(columns, start=1)
            if self._is_category_column(column)
        }

    def _is_email_column(self, column_name: str) -> bool:
        normalized_column = self._normalize_text(column_name)
        return any(word and word in normalized_column for word in self._email_column_words)

    def _is_category_column(self, column_name: str) -> bool:
        normalized_column = self._normalize_text(column_name)
        return any(word and word in normalized_column for word in self._category_column_words)

    def _category_fill(self, value: str) -> PatternFill:
        normalized_value = self._normalize_text(value)
        digest = hashlib.sha1(normalized_value.encode("utf-8")).hexdigest()
        color_index = int(digest[:8], 16) % len(self._category_fills)
        return self._category_fills[color_index]

    def _display_width(self, value: str) -> int:
        lines = [line.strip() for line in str(value).splitlines()] or [""]
        return max(len(line) for line in lines)

    def _normalize_text(self, value: str) -> str:
        normalized = normalize("NFD", str(value).casefold())
        return "".join(character for character in normalized if not combining(character))

    def export_name_checklist_pdf(
        self,
        output_path: str,
        guests: Iterable[GuestRecord],
        title: str,
    ) -> int:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        names = [self._first_line(guest.data.get("Nome", "")) for guest in guests]
        names = [name for name in names if name]

        self._ensure_gui_application()
        writer = QPdfWriter(str(path))
        writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        writer.setPageMargins(QMarginsF(14, 14, 14, 14), QPageLayout.Unit.Millimeter)
        writer.setResolution(96)

        painter = QPainter(writer)
        try:
            self._draw_name_checklist_pdf(painter, writer, title, names)
        finally:
            painter.end()
        return len(names)

    def _draw_name_checklist_pdf(
        self,
        painter: QPainter,
        writer: QPdfWriter,
        title: str,
        names: Sequence[str],
    ) -> None:
        margin = 46
        page_width = writer.width()
        page_height = writer.height()
        left = margin
        right = page_width - margin
        top = margin
        bottom = page_height - margin
        content_width = right - left

        title_font = QFont("Arial", 15)
        title_font.setBold(True)
        meta_font = QFont("Arial", 9)
        name_font = QFont("Arial", 11)
        line_pen = QPen(QColor("#d1d5db"))
        text_color = QColor("#111827")
        muted_color = QColor("#4b5563")

        page_number = 1

        def draw_header() -> float:
            painter.setPen(text_color)
            painter.setFont(title_font)
            painter.drawText(QRectF(left, top, content_width, 24), Qt.AlignmentFlag.AlignLeft, title)
            painter.setPen(muted_color)
            painter.setFont(meta_font)
            painter.drawText(
                QRectF(left, top + 28, content_width, 18),
                Qt.AlignmentFlag.AlignLeft,
                f"{len(names)} nomes para conferência",
            )
            painter.setPen(line_pen)
            painter.drawLine(left, top + 54, right, top + 54)
            return top + 74

        def draw_footer() -> None:
            painter.setPen(muted_color)
            painter.setFont(meta_font)
            painter.drawText(
                QRectF(left, bottom - 18, content_width, 16),
                Qt.AlignmentFlag.AlignRight,
                f"Página {page_number}",
            )

        y = draw_header()
        row_height = 30
        checkbox_size = 12
        checkbox_x = left
        name_x = checkbox_x + checkbox_size + 12
        name_width = content_width - checkbox_size - 12

        painter.setFont(name_font)
        metrics = QFontMetrics(name_font)

        for name in names:
            if y + row_height > bottom - 24:
                draw_footer()
                writer.newPage()
                page_number += 1
                y = draw_header()
                painter.setFont(name_font)
                metrics = QFontMetrics(name_font)

            checkbox_y = y + (row_height - checkbox_size) / 2
            painter.setPen(QPen(QColor("#111827")))
            painter.drawRect(QRectF(checkbox_x, checkbox_y, checkbox_size, checkbox_size))

            painter.setPen(text_color)
            display_name = metrics.elidedText(name, Qt.TextElideMode.ElideRight, int(name_width))
            painter.drawText(
                QRectF(name_x, y, name_width, row_height),
                Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft,
                display_name,
            )
            painter.setPen(line_pen)
            painter.drawLine(left, y + row_height - 1, right, y + row_height - 1)
            y += row_height

        draw_footer()

    def _ensure_gui_application(self) -> None:
        if QGuiApplication.instance() is None:
            self.__class__._qt_app = QGuiApplication([])

    def _first_line(self, value: object) -> str:
        for item in str(value or "").splitlines():
            clean_item = item.strip()
            if clean_item:
                return clean_item
        return ""

    def _category_value(self, value: object) -> str:
        for item in str(value or "").splitlines():
            clean_item = item.strip()
            if clean_item:
                return clean_item
        return "Sem categoria"

    def _safe_sheet_name(self, value: str, used_names: set[str]) -> str:
        base_name = re.sub(r"[\[\]:*?/\\]+", " ", str(value).strip())
        base_name = re.sub(r"\s{2,}", " ", base_name).strip("' ")
        if not base_name:
            base_name = "Planilha"
        base_name = base_name[:31]

        candidate = base_name
        counter = 2
        while candidate in used_names:
            suffix = f" {counter}"
            candidate = f"{base_name[:31 - len(suffix)]}{suffix}"
            counter += 1
        used_names.add(candidate)
        return candidate
