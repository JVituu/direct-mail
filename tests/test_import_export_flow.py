from collections.abc import Iterable
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from app.application.use_cases.consolidate_imported_data import ConsolidateImportedDataUseCase
from app.application.use_cases.create_guest_column import CreateGuestColumnUseCase
from app.application.use_cases.export_selected_guests import ExportSelectedGuestsUseCase
from app.application.use_cases.import_spreadsheet import ImportSpreadsheetUseCase
from app.application.use_cases.list_guest_filter_options import ListGuestFilterOptionsUseCase
from app.application.use_cases.list_guests import ListGuestsUseCase
from app.application.use_cases.manage_automatic_sheet import ManageAutomaticSheetUseCase
from app.application.use_cases.merge_workbooks import MergeWorkbooksUseCase
from app.application.use_cases.register_guest import RegisterGuestUseCase
from app.application.use_cases.review_duplicate_selection import ReviewDuplicateSelectionUseCase
from app.application.use_cases.update_guest_data import UpdateGuestDataUseCase
from app.application.use_cases.update_guest_selection import UpdateGuestSelectionUseCase
from app.domain.entities.guest_record import GuestRecord
from app.infrastructure.repositories.sqlite_guest_repository import SqliteGuestRepository
from app.infrastructure.spreadsheet.openpyxl_exporter import OpenpyxlSelectedGuestsExporter
from app.infrastructure.spreadsheet.openpyxl_reader import OpenpyxlSpreadsheetReader
from app.shared.utils.search_filter import encode_or_filter


class CapturingPdfExporter(OpenpyxlSelectedGuestsExporter):
    def __init__(self) -> None:
        self.pdf_names: list[str] = []

    def export_name_checklist_pdf(
        self,
        output_path: str,
        guests: Iterable[GuestRecord],
        title: str,
    ) -> int:
        self.pdf_names = [str(guest.data.get("Nome", "")) for guest in guests]
        Path(output_path).write_bytes(b"%PDF-1.4\n% test\n")
        return len([name for name in self.pdf_names if name])


class ImportExportFlowTest(unittest.TestCase):
    def test_consolidates_workbooks_and_keeps_most_complete_duplicate(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            first_spreadsheet = temp_path / "lista_original.xlsx"
            second_spreadsheet = temp_path / "lista_nova.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_contact_workbook(
                first_spreadsheet,
                "Amigos",
                ["Nome", "Telefone", "E-mail", "CEP", "ENDERECO"],
                [["Ana Silva", "3268-1279", "", "", ""]],
            )
            self._create_contact_workbook(
                second_spreadsheet,
                "Convidados",
                ["Nome", "Telefone", "E-mail", "CEP", "ENDERECO"],
                [["Ana da Silva", "", "ana@example.com", "50000-000", "Rua Alpha, 100"]],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_use_case = ImportSpreadsheetUseCase(repository, OpenpyxlSpreadsheetReader())

            first_import = import_use_case.execute_workbook(str(first_spreadsheet))
            import_use_case.execute_workbook(str(second_spreadsheet))

            result = ConsolidateImportedDataUseCase(repository).execute()

            self.assertEqual(result.workbook_id, first_import.workbook_id)
            self.assertEqual(result.moved_sheets, 1)
            self.assertEqual(result.removed_duplicates, 1)
            self.assertEqual(result.total_rows, 1)
            self.assertEqual(len(repository.list_workbooks()), 1)

            page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=first_import.workbook_id,
                page=0,
                page_size=50,
            )
            self.assertEqual(page.total_rows, 1)
            self.assertEqual(page.rows[0].data["Nome"], "Ana da Silva")
            self.assertEqual(page.rows[0].data["Telefone"], "3268-1279")
            self.assertEqual(page.rows[0].data["E-mail"], "ana@example.com")
            self.assertEqual(page.rows[0].data["CEP"], "50000-000")
            self.assertEqual(page.rows[0].data["ENDERECO"], "Rua Alpha, 100")

            duplicates_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=None,
                page=0,
                page_size=50,
                duplicates_only=True,
            )
            self.assertEqual(duplicates_page.total_rows, 0)

    def test_edits_blank_cells_when_column_is_missing_from_row_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "listas_com_colunas_diferentes.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            workbook = Workbook()
            first_sheet = workbook.active
            first_sheet.title = "Amigos"
            first_sheet.append(["Nome", "Telefone"])
            first_sheet.append(["Ana Silva", ""])
            second_sheet = workbook.create_sheet("Arquitetos")
            second_sheet.append(["Nome", "Endereço"])
            second_sheet.append(["Bruno Costa", "Rua Beta, 200"])
            workbook.save(spreadsheet_path)
            workbook.close()

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                repository,
                OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            unified_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            self.assertEqual(unified_page.columns, ("Lista", "Nome", "Endereço"))
            ana_row = next(row for row in unified_page.rows if row.data["Nome"] == "Ana Silva")
            self.assertEqual(ana_row.data["Endereço"], "")

            selection_use_case = UpdateGuestSelectionUseCase(repository)
            selection_use_case.set_guest_selected(ana_row.id, True)

            update_data_use_case = UpdateGuestDataUseCase(repository)
            update_data_use_case.execute(
                guest_id=ana_row.id,
                column_name="Endereço",
                value="Rua da Lista Final, 10",
                automatic=True,
            )

            selected_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(selected_page.rows[0].data["Endereço"], "Rua da Lista Final, 10")

            source_page_before_edit = ListGuestsUseCase(repository).execute(
                import_id=import_result.imported_sheets[0].import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            self.assertNotIn("Endereço", source_page_before_edit.columns)

            update_data_use_case.execute(
                guest_id=ana_row.id,
                column_name="Endereço",
                value="Rua da Origem, 20",
            )

            source_page_after_edit = ListGuestsUseCase(repository).execute(
                import_id=import_result.imported_sheets[0].import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            self.assertEqual(source_page_after_edit.columns, ("Nome", "Endereço"))
            self.assertEqual(source_page_after_edit.rows[0].data["Endereço"], "Rua da Origem, 20")

            selected_page_after_source_edit = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(
                selected_page_after_source_edit.rows[0].data["Endereço"],
                "Rua da Lista Final, 10",
            )

    def test_updates_automatic_guest_invitation_status(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "convites.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["Nome", "Telefone"],
                [["Ana Silva", "90000-0001"]],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                repository,
                OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            page = ListGuestsUseCase(repository).execute(
                import_id=import_result.imported_sheets[0].import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            selection_use_case = UpdateGuestSelectionUseCase(repository)
            selection_use_case.set_guest_selected(page.rows[0].id, True)

            automatic_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(automatic_page.rows[0].invitation_status, "pending")

            selection_use_case.set_automatic_guest_status(page.rows[0].id, "sent")
            sent_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(sent_page.rows[0].invitation_status, "sent")

            selection_use_case.set_automatic_guest_status(page.rows[0].id, "waiting")
            waiting_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(waiting_page.rows[0].invitation_status, "waiting")

            selection_use_case.set_automatic_guest_status(page.rows[0].id, "not_sent")
            not_sent_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(not_sent_page.rows[0].invitation_status, "not_sent")

            selection_use_case.set_automatic_contact_status(page.rows[0].id, "call", "done")
            selection_use_case.set_automatic_contact_status(page.rows[0].id, "email", "done")
            contact_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(
                contact_page.rows[0].contact_statuses,
                {"phone": "done", "mobile": "done", "email": "done"},
            )

    def test_lists_filter_options_and_filters_by_combined_terms(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "filtros.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["Nome", "Categoria", "ESTADO"],
                [
                    ["Ana Silva", "AMIGOS", "RECIFE - PE"],
                    ["Bruno Lima", "ADVOGADOS", "OLINDA - PE"],
                    ["Carla Souza", "AMIGOS", "JABOATAO - PE"],
                ],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                repository,
                OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            options = ListGuestFilterOptionsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
            )
            self.assertEqual(options.categories, ("ADVOGADOS", "AMIGOS"))
            self.assertEqual(options.locations, ("JABOATAO - PE", "OLINDA - PE", "RECIFE - PE"))

            filtered_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                search="AMIGOS RECIFE PE",
            )
            self.assertEqual(filtered_page.total_rows, 1)
            self.assertEqual(filtered_page.rows[0].data["Nome"], "Ana Silva")

            multi_category_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                search=encode_or_filter(["AMIGOS", "ADVOGADOS"]),
            )
            self.assertEqual(multi_category_page.total_rows, 3)

    def test_lists_only_one_category_column_when_sheet_has_duplicate_categories(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "categorias_duplicadas.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["Nome", "Categoria", "Categoria", "Telefone"],
                [
                    ["Ana Silva", "", "AMIGOS", "90000-0001"],
                    ["Bruno Lima", "ADVOGADOS", "", "90000-0002"],
                    ["Carla Souza", "AMIGOS", "VIP", "90000-0003"],
                ],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                repository,
                OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )

            self.assertEqual(page.columns, ("Lista", "Nome", "Categoria", "Telefone"))
            self.assertEqual(page.rows[0].data["Categoria"], "AMIGOS")
            self.assertEqual(page.rows[1].data["Categoria"], "ADVOGADOS")
            self.assertEqual(page.rows[2].data["Categoria"], "AMIGOS\nVIP")

    def test_hides_columns_without_values_from_guest_list(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "colunas_vazias.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["Nome", "Telefone", "Celular", "Contato - Atualizado", "Observação vazia", "Categoria"],
                [
                    ["Ana Silva", "3268-1279", "9", "", "", "AMIGOS"],
                    ["Bruno Lima", "3268-1280", "", "SIM", "", "ADVOGADOS"],
                ],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                repository,
                OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))
            imported_file = repository.get_import(import_result.imported_sheets[0].import_id)

            self.assertIsNotNone(imported_file)
            self.assertEqual(imported_file.columns, ("Nome", "Telefone", "Categoria"))

            page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )

            self.assertEqual(page.columns, ("Lista", "Nome", "Telefone", "Categoria"))
            self.assertNotIn("Celular", page.rows[0].data)
            self.assertNotIn("Contato - Atualizado", page.rows[0].data)
            self.assertNotIn("Observação vazia", page.rows[0].data)

    def test_creates_guest_column_for_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "nova_coluna.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["Nome", "Telefone"],
                [["Ana Silva", "90000-0001"]],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                repository,
                OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            updated_lists = CreateGuestColumnUseCase(repository).execute(
                "Observação",
                workbook_id=import_result.workbook_id,
            )

            self.assertEqual(updated_lists, 1)
            self.assertIn("Observação", repository.get_columns(None, import_result.workbook_id))

    def test_registers_manual_guest_in_unified_list(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "base.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["NOME", "Telefone", "E-mail", "ENDERECO", "Categoria"],
                [["Ana Silva", "90000-0001", "ana@example.com", "Rua A", "AMIGOS"]],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                repository,
                OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            result = RegisterGuestUseCase(repository).execute(
                {
                    "NOME": "Bruno Lima",
                    "Telefone": "3268-1279",
                    "Celular": "99999-0000",
                    "E-mail": "bruno@example.com",
                    "ENDEREÇO": "Rua B",
                    "Categoria": "ADVOGADOS",
                    "OBS": "Cadastro feito pelo sistema",
                },
                workbook_id=import_result.workbook_id,
            )

            self.assertEqual(result.workbook_id, import_result.workbook_id)
            self.assertEqual(result.sheet_name, "Cadastros manuais")

            page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            self.assertEqual(page.total_rows, 2)
            manual_row = next(row for row in page.rows if row.data.get("NOME") == "Bruno Lima")
            self.assertEqual(manual_row.data["Categoria"], "ADVOGADOS")
            self.assertEqual(manual_row.data["Telefone"], "3268-1279")

            options = ListGuestFilterOptionsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
            )
            self.assertEqual(options.categories, ("ADVOGADOS", "AMIGOS"))

    def test_merges_workbooks_and_preserves_tabs_selection_and_automatic_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            first_spreadsheet = temp_path / "lista_a.xlsx"
            second_spreadsheet = temp_path / "lista_b.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_multi_sheet_contact_workbook(first_spreadsheet, ["Amigos", "Arquitetos"])
            self._create_multi_sheet_contact_workbook(second_spreadsheet, ["Amigos", "Politicos"])

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_use_case = ImportSpreadsheetUseCase(repository, OpenpyxlSpreadsheetReader())

            first_import = import_use_case.execute_workbook(str(first_spreadsheet))
            second_import = import_use_case.execute_workbook(str(second_spreadsheet))

            second_amigos_page = ListGuestsUseCase(repository).execute(
                import_id=second_import.imported_sheets[0].import_id,
                workbook_id=second_import.workbook_id,
                page=0,
                page_size=50,
            )
            UpdateGuestSelectionUseCase(repository).set_guest_selected(
                guest_id=second_amigos_page.rows[0].id,
                selected=True,
            )

            result = MergeWorkbooksUseCase(repository).execute(
                source_workbook_id=second_import.workbook_id,
                target_workbook_id=first_import.workbook_id,
            )

            self.assertEqual(result.moved_sheets, 2)
            self.assertEqual(result.moved_rows, 2)
            self.assertEqual(result.target_total_rows, 4)

            workbooks = repository.list_workbooks()
            self.assertEqual(len(workbooks), 1)
            self.assertEqual(workbooks[0].id, first_import.workbook_id)
            self.assertEqual(workbooks[0].total_rows, 4)

            merged_imports = repository.list_imports(first_import.workbook_id)
            merged_sheet_names = [imported_sheet.sheet_name for imported_sheet in merged_imports]
            self.assertEqual(len(merged_imports), 4)
            self.assertIn("Amigos", merged_sheet_names)
            self.assertIn("Arquitetos", merged_sheet_names)
            self.assertIn("Politicos", merged_sheet_names)
            self.assertTrue(
                any(sheet_name.startswith("Amigos - lista_b") for sheet_name in merged_sheet_names)
            )

            moved_import = repository.get_import(second_import.imported_sheets[0].import_id)
            self.assertIsNotNone(moved_import)
            self.assertEqual(moved_import.workbook_id, first_import.workbook_id)

            automatic_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=first_import.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(automatic_page.total_rows, 1)
            self.assertTrue(automatic_page.rows[0].data["Lista"].startswith("Amigos - lista_b"))
            self.assertEqual(automatic_page.rows[0].data["Nome"], "Amigos 1")

            old_source_automatic_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=second_import.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(old_source_automatic_page.total_rows, 0)

    def test_detects_duplicates_across_workbooks_by_normalized_name(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            first_spreadsheet = temp_path / "lista_a.xlsx"
            second_spreadsheet = temp_path / "lista_b.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_single_contact_workbook(first_spreadsheet, "Amigos", "José da Silva", "(00) 90000-0000")
            self._create_single_contact_workbook(second_spreadsheet, "Políticos", "Jose Silva", "(00) 91111-1111")

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            reader = OpenpyxlSpreadsheetReader()
            import_use_case = ImportSpreadsheetUseCase(repository, reader)

            first_import = import_use_case.execute_workbook(str(first_spreadsheet))
            second_import = import_use_case.execute_workbook(str(second_spreadsheet))

            duplicates_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=None,
                page=0,
                page_size=50,
                duplicates_only=True,
            )

            self.assertEqual(duplicates_page.total_rows, 2)
            self.assertEqual({row.duplicate_reason for row in duplicates_page.rows}, {"Nome"})
            self.assertEqual({row.duplicate_count for row in duplicates_page.rows}, {2})

            first_page = ListGuestsUseCase(repository).execute(
                import_id=first_import.imported_sheets[0].import_id,
                workbook_id=first_import.workbook_id,
                page=0,
                page_size=50,
            )
            second_page = ListGuestsUseCase(repository).execute(
                import_id=second_import.imported_sheets[0].import_id,
                workbook_id=second_import.workbook_id,
                page=0,
                page_size=50,
            )

            self.assertEqual(first_page.rows[0].duplicate_reason, "Nome")
            self.assertEqual(first_page.rows[0].duplicate_count, 2)
            self.assertEqual(second_page.rows[0].duplicate_reason, "Nome")
            self.assertEqual(second_page.rows[0].duplicate_count, 2)

            review_use_case = ReviewDuplicateSelectionUseCase(repository)
            candidates = review_use_case.list_duplicate_candidates(first_page.rows[0].id)
            self.assertEqual({candidate.id for candidate in candidates}, {first_page.rows[0].id, second_page.rows[0].id})

            UpdateGuestSelectionUseCase(repository).set_guest_selected(first_page.rows[0].id, True)
            automatic_conflicts = review_use_case.list_automatic_conflicts(second_page.rows[0].id)
            self.assertEqual([conflict.id for conflict in automatic_conflicts], [first_page.rows[0].id])

    def test_ignores_weak_single_token_duplicate_names(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            first_spreadsheet = temp_path / "lista_a.xlsx"
            second_spreadsheet = temp_path / "lista_b.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            self._create_single_contact_workbook(first_spreadsheet, "Lista A", "FILHA", "(00) 90000-0000")
            self._create_single_contact_workbook(second_spreadsheet, "Lista B", "FILHA", "(00) 91111-1111")

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_use_case = ImportSpreadsheetUseCase(repository, OpenpyxlSpreadsheetReader())

            first_import = import_use_case.execute_workbook(str(first_spreadsheet))
            import_use_case.execute_workbook(str(second_spreadsheet))

            duplicates_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=None,
                page=0,
                page_size=50,
                duplicates_only=True,
            )
            first_page = ListGuestsUseCase(repository).execute(
                import_id=first_import.imported_sheets[0].import_id,
                workbook_id=first_import.workbook_id,
                page=0,
                page_size=50,
            )
            candidates = ReviewDuplicateSelectionUseCase(repository).list_duplicate_candidates(first_page.rows[0].id)

            self.assertEqual(duplicates_page.total_rows, 0)
            self.assertLessEqual(len(candidates), 1)

    def test_cleans_contact_values_imported_in_wrong_columns(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "contatos_baguncados.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Contatos"
            worksheet.append(["NOME", "Telefone", "Celular", "E-mail", "CEP", "ENDERECO", "EDIFICIO", "ESTADO"])
            worksheet.append([
                "Ana",
                "ana.extra@hotmail.com",
                "3333-4444",
                "ana@email.com",
                "81999998888",
                "",
                "",
                "",
            ])
            worksheet.append([
                "Bruno",
                "99999-8888",
                "bruno@example.com",
                "",
                "50000-000",
                "Rua Alpha, 100",
                "EDF TESTE",
                "PE",
            ])
            worksheet.append([
                "Carla",
                "Rua Beta, 200",
                "EDF CENTRAL",
                "",
                "",
                "",
                "",
                "",
            ])
            worksheet.append([
                "Diego",
                "32681279",
                "999998888",
                "",
                "51030000",
                "",
                "",
                "",
            ])
            worksheet.append([
                "Elaine",
                "",
                "",
                "",
                "51020-210\n33270845\n99119444",
                "",
                "",
                "",
            ])
            workbook.save(spreadsheet_path)
            workbook.close()

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                guest_repository=repository,
                spreadsheet_reader=OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            page = ListGuestsUseCase(repository).execute(
                import_id=import_result.imported_sheets[0].import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )

            ana = page.rows[0].data
            bruno = page.rows[1].data
            carla = page.rows[2].data
            diego = page.rows[3].data
            elaine = page.rows[4].data

            self.assertEqual(ana["Telefone"], "3333-4444")
            self.assertEqual(ana["Celular"], "81999998888")
            self.assertEqual(ana["E-mail"], "ana@email.com\nana.extra@hotmail.com")
            self.assertEqual(ana["CEP"], "")

            self.assertEqual(bruno["Telefone"], "")
            self.assertEqual(bruno["Celular"], "99999-8888")
            self.assertEqual(bruno["E-mail"], "bruno@example.com")
            self.assertEqual(bruno["CEP"], "50000-000")

            self.assertEqual(carla["Telefone"], "")
            self.assertEqual(carla["Celular"], "")
            self.assertEqual(carla["ENDERECO"], "Rua Beta, 200")
            self.assertEqual(carla["EDIFICIO"], "EDF CENTRAL")

            self.assertEqual(diego["Telefone"], "32681279")
            self.assertEqual(diego["Celular"], "999998888")
            self.assertEqual(diego["CEP"], "51030000")

            self.assertEqual(elaine["Telefone"], "33270845")
            self.assertEqual(elaine["Celular"], "99119444")
            self.assertEqual(elaine["CEP"], "51020-210")

            UpdateGuestSelectionUseCase(repository).set_guest_selected(page.rows[0].id, True)
            automatic_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            automatic_ana = automatic_page.rows[0].data
            self.assertEqual(automatic_ana["Telefone"], "3333-4444")
            self.assertEqual(automatic_ana["Celular"], "81999998888")
            self.assertEqual(automatic_ana["E-mail"], "ana@email.com\nana.extra@hotmail.com")

    def test_imports_simple_single_column_table(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "convidados.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Lista"
            worksheet.append(["Convidado"])
            worksheet.append(["Jose"])
            worksheet.append(["Maria"])
            workbook.save(spreadsheet_path)
            workbook.close()

            repository = SqliteGuestRepository(database_path)
            repository.initialize()

            import_result = ImportSpreadsheetUseCase(
                guest_repository=repository,
                spreadsheet_reader=OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            self.assertEqual(import_result.total_rows, 2)
            self.assertEqual(len(import_result.imported_sheets), 1)
            self.assertTrue(import_result.imported_sheets[0].is_selectable)

            page = ListGuestsUseCase(repository).execute(
                import_id=import_result.imported_sheets[0].import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )

            self.assertEqual(page.columns, ("Convidado",))
            self.assertEqual(len(page.rows[0].verification_code), 8)
            self.assertTrue(page.rows[0].verification_code.isdigit())
            self.assertNotEqual(page.rows[0].verification_code, page.rows[1].verification_code)
            self.assertEqual(page.rows[0].data["Convidado"], "Jose")

    def test_import_workbook_select_and_export_automatic_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "mala_direta.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"
            export_path = temp_path / "selecionados.xlsx"
            preserved_export_path = temp_path / "selecionados_preservados.xlsx"

            self._create_workbook(spreadsheet_path)

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            reader = OpenpyxlSpreadsheetReader()
            exporter = OpenpyxlSelectedGuestsExporter()

            import_result = ImportSpreadsheetUseCase(
                guest_repository=repository,
                spreadsheet_reader=reader,
                batch_size=2,
            ).execute_workbook(str(spreadsheet_path))

            self.assertEqual(import_result.total_rows, 6)
            self.assertEqual(len(import_result.imported_sheets), 3)
            self.assertEqual(import_result.skipped_sheets, tuple())
            self.assertEqual(len(repository.list_workbooks()), 1)
            repository.rename_workbook(import_result.workbook_id, "Agenda de teste")
            self.assertEqual(repository.list_workbooks()[0].display_name, "Agenda de teste")

            summary_sheet = self._find_sheet(import_result.imported_sheets, "Resumo")
            amigos_sheet = self._find_sheet(import_result.imported_sheets, "Amigos")
            politicos_sheet = self._find_sheet(import_result.imported_sheets, "Políticos")

            self.assertFalse(summary_sheet.is_selectable)
            self.assertTrue(amigos_sheet.is_selectable)
            self.assertTrue(politicos_sheet.is_selectable)

            summary_page = ListGuestsUseCase(repository).execute(
                import_id=summary_sheet.import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )

            self.assertEqual(summary_page.total_rows, 2)
            self.assertEqual(summary_page.columns, ("Categoria", "Quantidade de contatos"))
            self.assertFalse(summary_page.rows[0].selectable)

            amigos_page = ListGuestsUseCase(repository).execute(
                import_id=amigos_sheet.import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )

            self.assertEqual(amigos_page.total_rows, 2)
            self.assertEqual(amigos_page.columns, ("Nome", "Endereço", "Telefone", "Obra / Universo"))
            self.assertEqual(amigos_page.rows[0].data["Nome"], "Chandler Bing")
            self.assertEqual(len(amigos_page.rows[0].verification_code), 8)
            self.assertNotEqual(amigos_page.rows[0].verification_code, amigos_page.rows[1].verification_code)
            self.assertTrue(amigos_page.rows[0].selectable)

            code_search_page = ListGuestsUseCase(repository).execute(
                import_id=amigos_sheet.import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                search=amigos_page.rows[0].verification_code,
            )
            self.assertEqual(code_search_page.total_rows, 1)
            self.assertEqual(code_search_page.rows[0].data["Nome"], "Chandler Bing")

            selection_use_case = UpdateGuestSelectionUseCase(repository)
            updated_rows = selection_use_case.set_all_filtered_selected(
                import_id=amigos_sheet.import_id,
                workbook_id=import_result.workbook_id,
                selected=True,
                search="Friends",
            )
            self.assertEqual(updated_rows, 2)

            cleared_rows = selection_use_case.set_all_filtered_selected(
                import_id=amigos_sheet.import_id,
                workbook_id=import_result.workbook_id,
                selected=False,
                search="Monica",
            )
            self.assertEqual(cleared_rows, 1)

            selection_use_case.set_guest_selected(
                guest_id=amigos_page.rows[1].id,
                selected=True,
            )

            politicos_page = ListGuestsUseCase(repository).execute(
                import_id=politicos_sheet.import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            selection_use_case.set_guest_selected(
                guest_id=politicos_page.rows[0].id,
                selected=True,
            )

            update_data_use_case = UpdateGuestDataUseCase(repository)
            update_data_use_case.execute(
                guest_id=amigos_page.rows[0].id,
                column_name="Telefone",
                value="(00) 99999-0001",
            )
            update_data_use_case.execute(
                guest_id=politicos_page.rows[0].id,
                column_name="Nome",
                value="Ada Byron",
            )

            refreshed_amigos_page = ListGuestsUseCase(repository).execute(
                import_id=amigos_sheet.import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            self.assertEqual(refreshed_amigos_page.rows[0].data["Telefone"], "(00) 99999-0001")

            selected_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )

            self.assertEqual(selected_page.total_rows, 3)
            self.assertEqual(selected_page.columns, ("Lista", "Nome", "Endereço", "Telefone", "Obra / Universo"))
            self.assertEqual(selected_page.editable_columns, ("Nome", "Endereço", "Telefone", "Obra / Universo"))
            self.assertEqual(selected_page.rows[0].data["Telefone"], "(00) 90000-0001")
            self.assertEqual(selected_page.rows[2].data["Nome"], "Ada Lovelace")
            self.assertEqual(selected_page.rows[0].verification_code, amigos_page.rows[0].verification_code)

            update_data_use_case.execute(
                guest_id=selected_page.rows[0].id,
                column_name="Nome",
                value="Chandler Export",
                automatic=True,
            )

            automatic_page_after_edit = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
                selected_only=True,
            )
            source_page_after_automatic_edit = ListGuestsUseCase(repository).execute(
                import_id=amigos_sheet.import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            self.assertEqual(automatic_page_after_edit.rows[0].data["Nome"], "Chandler Export")
            self.assertEqual(source_page_after_automatic_edit.rows[0].data["Nome"], "Chandler Bing")

            automatic_sheet_use_case = ManageAutomaticSheetUseCase(repository)
            self.assertEqual(automatic_sheet_use_case.get_name(), "Planilha automática")
            automatic_sheet_use_case.rename("Lista final")
            self.assertEqual(automatic_sheet_use_case.get_name(), "Lista final")

            export_result = ExportSelectedGuestsUseCase(
                guest_repository=repository,
                exporter=exporter,
            ).execute(None, str(export_path), workbook_id=import_result.workbook_id)

            self.assertEqual(export_result.total_rows, 3)
            self.assertTrue(export_path.exists())
            self._assert_exported_content(export_path)

            repository.delete_workbook(import_result.workbook_id)
            self.assertEqual(repository.list_workbooks(), [])
            self.assertEqual(repository.list_imports(import_result.workbook_id), [])

            preserved_page = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=None,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(preserved_page.total_rows, 3)
            self.assertEqual(preserved_page.columns, selected_page.columns)
            self.assertEqual(preserved_page.rows[0].data["Nome"], "Chandler Export")

            preserved_export_result = ExportSelectedGuestsUseCase(
                guest_repository=repository,
                exporter=exporter,
            ).execute(None, str(preserved_export_path))

            self.assertEqual(preserved_export_result.total_rows, 3)
            self.assertTrue(preserved_export_path.exists())
            self._assert_exported_content(preserved_export_path)

            cleared_rows = automatic_sheet_use_case.clear()
            self.assertEqual(cleared_rows, 3)
            self.assertEqual(automatic_sheet_use_case.get_name(), "Planilha automática")

            selected_page_after_clear = ListGuestsUseCase(repository).execute(
                import_id=None,
                workbook_id=None,
                page=0,
                page_size=50,
                selected_only=True,
            )
            self.assertEqual(selected_page_after_clear.total_rows, 0)

    def test_exports_selected_guests_as_names_and_by_category(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "categorias.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"
            names_export_path = temp_path / "nomes.xlsx"
            category_export_path = temp_path / "categorias_exportadas.xlsx"
            unified_export_path = temp_path / "lista_unificada.xlsx"
            final_pdf_export_path = temp_path / "lista_final_presenca.pdf"
            unified_pdf_export_path = temp_path / "lista_unificada_presenca.pdf"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["Nome", "NOME", "Categoria", "Telefone"],
                [
                    ["", "Ana Silva", "Amigos", "90000-0001"],
                    ["", "Bruno Lima", "Advogados", "90000-0002"],
                    ["", "Carla Souza", "Amigos", "90000-0003"],
                ],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                guest_repository=repository,
                spreadsheet_reader=OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            page = ListGuestsUseCase(repository).execute(
                import_id=import_result.imported_sheets[0].import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            UpdateGuestSelectionUseCase(repository).set_page_selected(
                [row.id for row in page.rows[:2]],
                True,
            )

            export_use_case = ExportSelectedGuestsUseCase(
                guest_repository=repository,
                exporter=OpenpyxlSelectedGuestsExporter(),
            )
            names_result = export_use_case.execute(
                None,
                str(names_export_path),
                workbook_id=import_result.workbook_id,
                export_mode="names",
            )
            category_result = export_use_case.execute(
                None,
                str(category_export_path),
                workbook_id=import_result.workbook_id,
                export_mode="category",
                export_scope="unified",
            )
            unified_result = export_use_case.execute(
                None,
                str(unified_export_path),
                workbook_id=import_result.workbook_id,
                export_mode="complete",
                export_scope="unified",
            )
            final_pdf_result = export_use_case.execute(
                None,
                str(final_pdf_export_path),
                workbook_id=import_result.workbook_id,
                export_format="pdf",
            )
            unified_pdf_result = export_use_case.execute(
                None,
                str(unified_pdf_export_path),
                workbook_id=import_result.workbook_id,
                export_scope="unified",
                export_format="pdf",
            )

            self.assertEqual(names_result.total_rows, 2)
            self.assertEqual(names_result.columns, ("Nome",))
            self.assertEqual(category_result.total_rows, 3)
            self.assertEqual(unified_result.total_rows, 3)
            self.assertEqual(unified_result.sheet_name, "Lista unificada")
            self.assertEqual(final_pdf_result.total_rows, 2)
            self.assertEqual(final_pdf_result.columns, ("Nome",))
            self.assertEqual(final_pdf_result.file_name, "lista_final_presenca.pdf")
            self.assertEqual(unified_pdf_result.total_rows, 3)
            self.assertEqual(unified_pdf_result.columns, ("Nome",))
            self.assertEqual(unified_pdf_result.file_name, "lista_unificada_presenca.pdf")
            self.assertGreater(final_pdf_export_path.stat().st_size, 1000)
            self.assertGreater(unified_pdf_export_path.stat().st_size, 1000)
            self.assertEqual(final_pdf_export_path.read_bytes()[:4], b"%PDF")
            self.assertEqual(unified_pdf_export_path.read_bytes()[:4], b"%PDF")

            names_workbook = load_workbook(names_export_path, read_only=True)
            try:
                names_rows = list(names_workbook["Nomes"].iter_rows(values_only=True))
            finally:
                names_workbook.close()
            self.assertEqual(
                names_rows,
                [
                    ("Nome",),
                    ("ANA SILVA",),
                    ("BRUNO LIMA",),
                ],
            )

            unified_workbook = load_workbook(unified_export_path)
            try:
                unified_worksheet = unified_workbook["Lista unificada"]
                unified_rows = list(unified_worksheet.iter_rows(values_only=True))
                self.assertEqual(unified_worksheet.freeze_panes, "A2")
                self.assertEqual(unified_worksheet.auto_filter.ref, "A1:D4")
                self.assertEqual(unified_worksheet["A1"].fill.fgColor.rgb, "00245783")
                self.assertGreater(unified_worksheet.column_dimensions["A"].width, 10)
            finally:
                unified_workbook.close()
            self.assertEqual(len(unified_rows), 4)
            self.assertEqual(unified_rows[0], ("Lista", "Nome", "Categoria", "Telefone"))
            self.assertEqual(unified_rows[3][1], "CARLA SOUZA")

            category_workbook = load_workbook(category_export_path, read_only=True)
            try:
                self.assertIn("Resumo", category_workbook.sheetnames)
                self.assertIn("Amigos", category_workbook.sheetnames)
                self.assertIn("Advogados", category_workbook.sheetnames)
                summary_rows = list(category_workbook["Resumo"].iter_rows(values_only=True))
                amigos_rows = list(category_workbook["Amigos"].iter_rows(values_only=True))
                advogados_rows = list(category_workbook["Advogados"].iter_rows(values_only=True))
            finally:
                category_workbook.close()

            self.assertEqual(summary_rows, [("Categoria", "Quantidade"), ("ADVOGADOS", 1), ("AMIGOS", 2)])
            self.assertEqual(amigos_rows[0], ("Lista", "Nome", "Categoria", "Telefone"))
            self.assertEqual(amigos_rows[1][1], "ANA SILVA")
            self.assertEqual(amigos_rows[2][1], "CARLA SOUZA")
            self.assertEqual(advogados_rows[1][1], "BRUNO LIMA")

    def test_exports_pdf_names_with_automatic_honorifics(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "tratamentos.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"
            pdf_export_path = temp_path / "lista_presenca.pdf"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["Nome", "Tratamento", "Sexo"],
                [
                    ["Ana Silva", "", ""],
                    ["Bruno Lima", "", ""],
                    ["Patricia Gomes", "SRA", ""],
                    ["Carlos Rocha", "SR", ""],
                    ["Alex Souza", "", "F"],
                    ["Taylor Costa", "", ""],
                ],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                guest_repository=repository,
                spreadsheet_reader=OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            page = ListGuestsUseCase(repository).execute(
                import_id=import_result.imported_sheets[0].import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            UpdateGuestSelectionUseCase(repository).set_page_selected(
                [row.id for row in page.rows],
                True,
            )

            exporter = CapturingPdfExporter()
            result = ExportSelectedGuestsUseCase(
                guest_repository=repository,
                exporter=exporter,
            ).execute(
                None,
                str(pdf_export_path),
                workbook_id=import_result.workbook_id,
                export_format="pdf",
            )

            self.assertEqual(result.total_rows, 6)
            self.assertEqual(
                exporter.pdf_names,
                [
                    "SRA. ANA SILVA",
                    "SR. BRUNO LIMA",
                    "SRA. PATRICIA GOMES",
                    "SR. CARLOS ROCHA",
                    "SRA. ALEX SOUZA",
                    "SR(A). TAYLOR COSTA",
                ],
            )

    def test_exports_pdf_institutions_with_representatives(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            spreadsheet_path = temp_path / "instituicoes.xlsx"
            database_path = temp_path / "mala_direta.sqlite3"
            pdf_export_path = temp_path / "lista_instituicoes.pdf"

            self._create_contact_workbook(
                spreadsheet_path,
                "Lista",
                ["Instituição", "Responsável", "Nome", "Sexo"],
                [
                    ["Instituto Ricardo Brennand", "Arnaldo Lima", "", "M"],
                    ["Empresa Y", "Morgana Costa", "", "F"],
                    ["", "", "Ana Silva", "F"],
                ],
            )

            repository = SqliteGuestRepository(database_path)
            repository.initialize()
            import_result = ImportSpreadsheetUseCase(
                guest_repository=repository,
                spreadsheet_reader=OpenpyxlSpreadsheetReader(),
            ).execute_workbook(str(spreadsheet_path))

            page = ListGuestsUseCase(repository).execute(
                import_id=import_result.imported_sheets[0].import_id,
                workbook_id=import_result.workbook_id,
                page=0,
                page_size=50,
            )
            UpdateGuestSelectionUseCase(repository).set_page_selected(
                [row.id for row in page.rows],
                True,
            )

            exporter = CapturingPdfExporter()
            result = ExportSelectedGuestsUseCase(
                guest_repository=repository,
                exporter=exporter,
            ).execute(
                None,
                str(pdf_export_path),
                workbook_id=import_result.workbook_id,
                export_format="pdf",
            )

            self.assertEqual(result.total_rows, 3)
            self.assertEqual(
                exporter.pdf_names,
                [
                    "INSTITUTO RICARDO BRENNAND - SR. ARNALDO LIMA",
                    "EMPRESA Y - SRA. MORGANA COSTA",
                    "SRA. ANA SILVA",
                ],
            )

    def _create_workbook(self, path: Path) -> None:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Resumo"
        summary.append(["Workbook criado para testes de mala direta", None, None, None])
        summary.merge_cells("A1:D1")
        summary.append(["Resumo das categorias importadas", None, None, None])
        summary.append(["Categoria", "Quantidade de contatos"])
        summary.append(["Amigos", 2])
        summary.append(["Políticos", 2])

        self._append_contact_sheet(
            workbook,
            "Amigos",
            [
                ["Chandler Bing", "Rua da Comédia, 101", "(00) 90000-0001", "Friends"],
                ["Monica Geller", "Rua do Café, 303", "(00) 90000-0003", "Friends"],
            ],
        )
        self._append_contact_sheet(
            workbook,
            "Políticos",
            [
                ["Ada Lovelace", "Rua Analítica, 1843", "(00) 91000-0001", "Computação"],
                ["Machado de Assis", "Rua Cosme Velho, 18", "(00) 91000-0002", "Literatura"],
            ],
        )

        workbook.save(path)
        workbook.close()

    def _create_single_contact_workbook(self, path: Path, sheet_name: str, name: str, phone: str) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.append(["Nome", "Telefone"])
        worksheet.append([name, phone])
        workbook.save(path)
        workbook.close()

    def _create_contact_workbook(
        self,
        path: Path,
        sheet_name: str,
        columns: list[str],
        rows: list[list[str]],
    ) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.append(columns)
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()

    def _create_multi_sheet_contact_workbook(self, path: Path, sheet_names: list[str]) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name in sheet_names:
            worksheet = workbook.create_sheet(sheet_name)
            worksheet.append(["Nome", "Telefone"])
            worksheet.append([f"{sheet_name} 1", "(00) 90000-0000"])
        workbook.save(path)
        workbook.close()

    def _append_contact_sheet(
        self,
        workbook: Workbook,
        sheet_name: str,
        rows: list[list[str]],
    ) -> None:
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append([f"Contatos fictícios — {sheet_name}", None, None, None])
        worksheet.merge_cells("A1:D1")
        worksheet.append(["Todos os nomes, endereços e telefones desta aba são fictícios.", None, None, None])
        worksheet.append(["Nome", "Endereço", "Telefone", "Obra / Universo"])
        for row in rows:
            worksheet.append(row)

    def _assert_exported_content(self, path: Path) -> None:
        workbook = load_workbook(path, read_only=True)
        try:
            worksheet = workbook["Selecionados"]
            rows = list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()

        self.assertEqual(rows[0], ("Lista", "Nome", "Endereço", "Telefone", "Obra / Universo"))
        self.assertEqual(rows[1], ("AMIGOS", "CHANDLER EXPORT", "RUA DA COMÉDIA, 101", "(00) 90000-0001", "FRIENDS"))
        self.assertEqual(rows[2], ("AMIGOS", "MONICA GELLER", "RUA DO CAFÉ, 303", "(00) 90000-0003", "FRIENDS"))
        self.assertEqual(rows[3], ("POLÍTICOS", "ADA LOVELACE", "RUA ANALÍTICA, 1843", "(00) 91000-0001", "COMPUTAÇÃO"))
        self.assertEqual(len(rows), 4)

    def _find_sheet(self, sheets: object, sheet_name: str) -> object:
        for sheet in sheets:
            if sheet.sheet_name == sheet_name:
                return sheet
        raise AssertionError(f"Aba não encontrada: {sheet_name}")


if __name__ == "__main__":
    unittest.main()
